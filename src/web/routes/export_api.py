"""Export API — SSE-streaming endpoints for chat export, wordcloud, report, employee."""
import hashlib
import os
import re
import sqlite3
import sys
import threading
from collections import defaultdict
from datetime import datetime
from flask import Blueprint, request, current_app, jsonify, send_from_directory

_BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _BASE not in sys.path:
    sys.path.insert(0, _BASE)

if getattr(sys, 'frozen', False):
    _DATA_ROOT = os.path.dirname(sys.executable)
else:
    _DATA_ROOT = os.path.normpath(os.path.join(_BASE, '..'))

from web.sse import create_sse_progress, sse_response

export_bp = Blueprint('export_api', __name__, url_prefix='/api/export')
_stdout_lock = threading.Lock()


def _decrypted_dir():
    return current_app.config.get('DECRYPTED_DIR', '')


def _scan_chats_flat(decrypted_dir: str, name_filter: str = ""):
    """Scan flat backup layout (no contact/session subdirs) for matching chats.

    Returns a list of chat_info dicts (compatible with scan_chats output) by
    reading Name2Id tables directly from message_*.db files in decrypted_dir.
    Each dict has: username, display_name, tables[{db_idx, db_path, table_name}].
    """
    if not os.path.isdir(decrypted_dir):
        return []

    # Find message DBs in flat layout (files directly in root dir)
    db_paths = []
    for f in os.listdir(decrypted_dir):
        if re.match(r'message_\d+\.db', f, re.IGNORECASE):
            db_paths.append(os.path.join(decrypted_dir, f))
    if not db_paths:
        return []
    db_paths.sort()

    # Build hash→username mapping from Name2Id tables
    hash_to_username = {}
    for db_path in db_paths:
        try:
            conn = sqlite3.connect(db_path)
            conn.text_factory = bytes
            for (uname,) in conn.execute("SELECT user_name FROM Name2Id"):
                if uname:
                    u = uname.decode("utf-8", errors="replace") if isinstance(uname, bytes) else str(uname)
                    if u:
                        hash_to_username[hashlib.md5(u.encode()).hexdigest()] = u
            conn.close()
        except sqlite3.Error:
            pass

    # Build chat list from Msg_ tables
    chat_map = defaultdict(lambda: {"tables": [], "total_msgs": 0})
    for db_path in db_paths:
        try:
            conn = sqlite3.connect(db_path)
            conn.text_factory = bytes
            tables = conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'Msg_%'"
            ).fetchall()
            for (tname,) in tables:
                tn = tname.decode("utf-8", errors="replace") if isinstance(tname, bytes) else str(tname)
                h = tn[4:]  # Strip "Msg_" prefix
                uname = hash_to_username.get(h, f"unknown_{h[:12]}")
                try:
                    count = conn.execute(f'SELECT COUNT(*) FROM [{tn}]').fetchone()[0]
                except sqlite3.Error:
                    count = 0
                chat_map[uname]["tables"].append({
                    "db_idx": 0,
                    "db_path": db_path,
                    "table_name": tn,
                })
                chat_map[uname]["total_msgs"] += count
            conn.close()
        except sqlite3.Error:
            pass

    # Load contacts from neighbouring output/decrypted for display name resolution
    # (flat backup layouts lack contact.db, but the old decrypted dir may have one)
    id_to_name = {}
    name_to_id = {}
    _project_root = _DATA_ROOT  # project root / exe dir
    contact_candidates = [
        os.path.join(decrypted_dir, "contact", "contact.db"),
        os.path.join(_project_root, "output", "decrypted", "contact", "contact.db"),
    ]
    for cc in contact_candidates:
        cc = os.path.normpath(cc)
        if os.path.exists(cc):
            try:
                cconn = sqlite3.connect(cc)
                cconn.text_factory = bytes
                for r in cconn.execute(
                    "SELECT id, username, remark, nick_name, alias FROM contact"
                ):
                    cid = r[0]
                    uname = (r[1] or b"").decode("utf-8", errors="replace").strip() if isinstance(r[1], bytes) else str(r[1] or "").strip()
                    remark = (r[2] or b"").decode("utf-8", errors="replace").strip() if isinstance(r[2], bytes) else str(r[2] or "").strip()
                    nick = (r[3] or b"").decode("utf-8", errors="replace").strip() if isinstance(r[3], bytes) else str(r[3] or "").strip()
                    alias = (r[4] or b"").decode("utf-8", errors="replace").strip() if isinstance(r[4], bytes) else str(r[4] or "").strip()
                    display = remark if (remark and remark != uname) else (nick if (nick and nick != uname) else (alias if (alias and alias != uname) else uname))
                    if cid and display:
                        id_to_name[cid] = display
                    if uname:
                        name_to_id[uname] = cid
                cconn.close()
            except sqlite3.Error:
                pass

    # Build result list, filter by name if provided
    result = []
    filter_lower = name_filter.lower() if name_filter else ""
    for uname, info in chat_map.items():
        # Resolve display name via contact db
        display = uname
        if uname in name_to_id:
            cid = name_to_id[uname]
            if cid in id_to_name:
                display = id_to_name[cid]
        if display == uname and uname.endswith("@chatroom") and len(uname) > 20:
            display = f"群聊({uname[:12]}...)"
        if filter_lower:
            if filter_lower not in uname.lower() and filter_lower not in display.lower():
                continue
        result.append({
            "username": uname,
            "display_name": display,
            "msg_count": info["total_msgs"],
            "tables": info["tables"],
        })

    result.sort(key=lambda x: x["msg_count"], reverse=True)
    return result


def _load_all_chats(decrypted_dir, own_wxid=None, prefer_fast=False, name_filter='',
                    with_counts=False):
    """加载会话列表（轻量优先，失败回退完整扫描 / 扁平布局）。"""
    all_chats = []
    if prefer_fast:
        try:
            from chatlab_pull_server import fast_chat_list
            all_chats = fast_chat_list(
                decrypted_dir, own_wxid=own_wxid, with_counts=with_counts) or []
        except Exception:
            all_chats = []
    if not all_chats:
        try:
            from chat_list import scan_chats
            all_chats, _, _ = scan_chats(decrypted_dir)
        except Exception:
            all_chats = []
    if not all_chats:
        all_chats = _scan_chats_flat(decrypted_dir, name_filter)
    return all_chats or []


def _resolve_targets(all_chats, contact='', contacts=None, display_name_hint=''):
    """按单选/多选解析目标会话。

    Returns: (targets, fuzzy_matches_or_None)
      - targets 非空：可直接导出
      - fuzzy_matches 非空：需前端二次选择
    """
    by_user = {c['username']: c for c in all_chats}
    selected_ids = []
    if contacts:
        if isinstance(contacts, str):
            contacts = [c.strip() for c in contacts.split(',') if c.strip()]
        selected_ids = [str(c).strip() for c in contacts if str(c).strip()]
    elif contact:
        selected_ids = [str(contact).strip()]

    if not selected_ids:
        return [], None

    # 多选：精确匹配 username / display_name
    if len(selected_ids) > 1 or (contacts and len(selected_ids) >= 1):
        targets = []
        missing = []
        for sid in selected_ids:
            hit = by_user.get(sid)
            if not hit:
                for c in all_chats:
                    if c.get('display_name') == sid:
                        hit = c
                        break
            if hit:
                targets.append(hit)
            else:
                missing.append(sid)
        if not targets:
            return [], [{'username': m, 'display_name': m, 'msg_count': 0} for m in missing]
        return targets, None

    # 单选：精确 → 模糊
    key = selected_ids[0]
    for c in all_chats:
        if c['username'] == key or c.get('display_name') == key:
            if display_name_hint and not c.get('display_name'):
                c = dict(c)
                c['display_name'] = display_name_hint
            return [c], None

    key_lower = key.lower()
    fuzzy_matches = []
    for c in all_chats:
        if (key_lower in (c.get('display_name') or '').lower()
                or key_lower in c['username'].lower()):
            fuzzy_matches.append({
                'username': c['username'],
                'display_name': c.get('display_name') or c['username'],
                'msg_count': c.get('msg_count', 0),
                'is_group': c['username'].endswith('@chatroom'),
            })
    if not fuzzy_matches:
        return [], None
    if len(fuzzy_matches) == 1:
        return [by_user[fuzzy_matches[0]['username']]], None
    return [], fuzzy_matches


@export_bp.route('/chats', methods=['GET'])
def list_export_chats():
    """GET /api/export/chats — 供导出页勾选会话（支持 q / kind=all|group|contact）。"""
    q = (request.args.get('q') or '').strip().lower()
    kind = (request.args.get('kind') or 'all').strip().lower()
    decrypted_dir = _decrypted_dir()
    own_wxid = current_app.config.get('WXID')
    try:
        # with_counts=True so the picker shows real message counts (not always 0)
        chats = _load_all_chats(
            decrypted_dir, own_wxid=own_wxid, prefer_fast=True,
            name_filter=q, with_counts=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e), 'chats': []}), 500

    result = []
    for c in chats:
        uname = c.get('username') or ''
        display = c.get('display_name') or uname
        is_group = uname.endswith('@chatroom')
        if kind == 'group' and not is_group:
            continue
        if kind in ('contact', 'private') and is_group:
            continue
        if q and q not in display.lower() and q not in uname.lower():
            continue
        result.append({
            'username': uname,
            'display_name': display,
            'msg_count': int(c.get('msg_count') or 0),
            'is_group': is_group,
        })
    result.sort(key=lambda x: (-x['msg_count'], x['display_name']))
    return jsonify({'ok': True, 'chats': result, 'total': len(result)})


@export_bp.route('/chat', methods=['POST'])
def export_chat():
    """POST /api/export/chat — Export chat messages.

    Body:
      contact / contacts: 单个或多个会话 username
      keyword: 关键词（空格=AND，| =OR）
      date_start / date_end: YYYY-MM-DD
      format: txt | html | xlsx | csv | md | json | pgsql | chatlab-jsonl | chatlab-json
      merge: 多会话 txt/html 是否合并（xlsx/csv 始终合并）
      include_asr: 是否带上已转写语音文字
      privacy_redact: 导出前脱敏昵称/手机号/身份证（分享演示）
      privacy: {nicknames, phones, id_cards, style: placeholder|asterisk}
    """
    data = request.get_json(silent=True) or {}
    contact = data.get('contact', '') or ''
    contacts = data.get('contacts') or data.get('usernames') or []
    display_name_hint = data.get('display_name', '')
    date_start = data.get('date_start', '')
    date_end = data.get('date_end', '')
    keyword = (data.get('keyword') or '').strip()
    fmt = (data.get('format') or 'txt').strip().lower()
    merge = data.get('merge')
    if merge is None:
        merge = fmt in ('xlsx', 'csv', 'html', 'json', 'pgsql', 'sql')
    from engine.services.privacy_redact import parse_privacy_options
    privacy = parse_privacy_options(data)

    decrypted_dir = _decrypted_dir()
    # 注意：后台线程没有 Flask 应用上下文，配置必须在此处（请求线程内）取出
    own_wxid = current_app.config.get('WXID')
    push, gen = create_sse_progress()

    def _run():
        try:
            from chat_export import export_chat as _export_one, export_chats
            from engine.constants import TZ

            is_chatlab = fmt.startswith('chatlab')
            push('export', '正在扫描聊天列表...', 0.1)
            # Prefer full scan for real export so every target has table_name + counts.
            # ChatLab discovery can stay on the fast path.
            all_chats = _load_all_chats(
                decrypted_dir, own_wxid=own_wxid,
                prefer_fast=is_chatlab,
                name_filter=contact if isinstance(contact, str) else '',
                with_counts=False)

            targets, fuzzy = _resolve_targets(
                all_chats, contact=contact, contacts=contacts,
                display_name_hint=display_name_hint)
            if fuzzy:
                push.select(fuzzy)
                return
            if not targets:
                tip = contact or (','.join(contacts) if contacts else '')
                push.error(f"未找到联系人/群聊: {tip or '（未选择）'}")
                return

            # Safety: if a fast-list target still lacks table_name, enrich from full scan
            need_enrich = [t for t in targets
                           if not any((x or {}).get('table_name')
                                      for x in (t.get('tables') or []))]
            if need_enrich:
                push('export', '补全会话表信息...', 0.15)
                full = _load_all_chats(
                    decrypted_dir, own_wxid=own_wxid, prefer_fast=False)
                by_user = {c['username']: c for c in full}
                fixed = []
                for t in targets:
                    full_t = by_user.get(t['username'])
                    fixed.append(full_t if full_t else t)
                targets = fixed

            start_ts = None
            end_ts = None
            if date_start:
                dt = datetime.strptime(date_start, '%Y-%m-%d').replace(tzinfo=TZ)
                start_ts = int(dt.timestamp())
            if date_end:
                dt = datetime.strptime(date_end, '%Y-%m-%d').replace(
                    hour=23, minute=59, second=59, tzinfo=TZ)
                end_ts = int(dt.timestamp())

            out_dir = os.path.join(_DATA_ROOT, 'export')
            os.makedirs(out_dir, exist_ok=True)

            def _print(msg):
                push('export', msg, 0.5)

            # 会话报告包：对话 HTML + 词云 + 统计 → ZIP
            if fmt in ('session-pack', 'session_pack', 'pack'):
                import json as _json
                import shutil
                import tempfile
                import zipfile as _zipfile
                from chat_export import export_chat as _export_one
                from wordcloud_gen import generate_wordcloud

                stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                out_dir = os.path.join(_DATA_ROOT, 'export')
                os.makedirs(out_dir, exist_ok=True)
                tmp_root = tempfile.mkdtemp(prefix='session_pack_')
                try:
                    pack_files = []  # (arcname, abs_path)
                    total_msgs = 0
                    for i, target in enumerate(targets):
                        raw_name = (target.get('display_name')
                                    or target.get('username') or 'chat')
                        if len(targets) == 1 and display_name_hint:
                            raw_name = display_name_hint
                        safe = re.sub(r'[\\/:*?"<>|\s]+', '_', str(raw_name)).strip('_.') or 'chat'
                        folder = safe if len(targets) > 1 else '.'
                        work = os.path.join(tmp_root, safe) if len(targets) > 1 else tmp_root
                        os.makedirs(work, exist_ok=True)

                        push('export', '导出对话 %s (%d/%d)...' % (
                            raw_name, i + 1, len(targets)),
                             0.2 + 0.5 * i / max(len(targets), 1))
                        count, chat_path = _export_one(
                            target, work,
                            start_ts=start_ts, end_ts=end_ts,
                            keyword=keyword or None,
                            print_fn=_print, fmt='html',
                            display_name=raw_name,
                            include_asr=bool(data.get('include_asr')),
                            privacy=privacy)
                        total_msgs += count or 0
                        if chat_path and os.path.isfile(chat_path):
                            dest = os.path.join(work, 'chat.html')
                            if os.path.abspath(chat_path) != os.path.abspath(dest):
                                shutil.copy2(chat_path, dest)
                            pack_files.append((
                                os.path.join(folder, 'chat.html').replace('\\', '/'),
                                dest))

                        push('export', '生成词云 %s...' % raw_name, 0.55)
                        wc_path = os.path.join(work, 'wordcloud.html')
                        try:
                            wc = generate_wordcloud(
                                decrypted_dir, chat_info=target,
                                out_path=wc_path,
                                start_ts=start_ts, end_ts=end_ts,
                                print_fn=_print)
                            if wc and os.path.isfile(wc):
                                pack_files.append((
                                    os.path.join(folder, 'wordcloud.html').replace('\\', '/'),
                                    wc))
                                # 离线图表脚本（与 HTML 同目录）
                                for asset in ('echarts.min.js',
                                              'echarts-wordcloud.min.js'):
                                    ap = os.path.join(os.path.dirname(wc), asset)
                                    if os.path.isfile(ap):
                                        pack_files.append((
                                            os.path.join(folder, asset).replace('\\', '/'),
                                            ap))
                        except Exception as wc_err:
                            _print('词云跳过：%s' % wc_err)

                        stats = {
                            'display_name': raw_name,
                            'username': target.get('username'),
                            'is_group': bool(
                                (target.get('username') or '').endswith('@chatroom')),
                            'msg_count_exported': count or 0,
                            'keyword': keyword or None,
                            'date_start': date_start or None,
                            'date_end': date_end or None,
                            'generated_at': stamp,
                        }
                        stats_path = os.path.join(work, 'stats.json')
                        with open(stats_path, 'w', encoding='utf-8') as sf:
                            _json.dump(stats, sf, ensure_ascii=False, indent=2)
                        pack_files.append((
                            os.path.join(folder, 'stats.json').replace('\\', '/'),
                            stats_path))

                        readme = os.path.join(work, 'README.txt')
                        with open(readme, 'w', encoding='utf-8') as rf:
                            rf.write(
                                'WeChat Export 会话报告包\n'
                                '会话：%s\n'
                                '导出消息：%d 条\n'
                                '包含：chat.html（对话）/ wordcloud.html（词云）/ stats.json（统计）\n'
                                % (raw_name, count or 0))
                        pack_files.append((
                            os.path.join(folder, 'README.txt').replace('\\', '/'),
                            readme))

                    zip_name = 'session_pack_%s.zip' % stamp
                    zip_path = os.path.join(out_dir, zip_name)
                    with _zipfile.ZipFile(zip_path, 'w', _zipfile.ZIP_DEFLATED) as zf:
                        seen = set()
                        for arc, path in pack_files:
                            if not path or not os.path.isfile(path):
                                continue
                            if arc in seen:
                                continue
                            seen.add(arc)
                            zf.write(path, arcname=arc)
                    push.done({
                        'msg_count': total_msgs,
                        'download_url': '/api/export/download/' + zip_name,
                        'filename': zip_name,
                        'note': '会话报告包：对话 HTML + 词云 + 统计 JSON',
                    })
                finally:
                    shutil.rmtree(tmp_root, ignore_errors=True)
                return

            # ChatLab：逐会话导出（多选时打包 zip）
            if is_chatlab:
                from chatlab_export import export_chatlab
                import zipfile as _zipfile
                sub_fmt = 'json' if fmt.endswith('json') else 'jsonl'
                files = []
                total = 0
                for i, target in enumerate(targets):
                    raw_name = (target.get('display_name')
                                or target.get('username') or 'chat')
                    if len(targets) == 1 and display_name_hint:
                        raw_name = display_name_hint
                    safe = re.sub(r'[\\/:*?"<>|\s]+', '_', str(raw_name)).strip('_.') or 'chat'
                    out_path = os.path.join(out_dir, 'chatlab_%s.%s' % (safe, sub_fmt))
                    push('export', 'ChatLab 导出 %s (%d/%d)...' % (
                        raw_name, i + 1, len(targets)), 0.2 + 0.6 * i / max(len(targets), 1))
                    res = export_chatlab(
                        decrypted_dir, target, out_path, fmt=sub_fmt,
                        own_wxid=own_wxid,
                        resume=bool(data.get('resume', True)),
                        start_ts=start_ts, end_ts=end_ts,
                        keyword=keyword or None,
                        print_fn=_print)
                    path = res.get('path') or out_path
                    files.append(path)
                    total += res.get('count_total') or 0
                if len(files) == 1:
                    fname = os.path.basename(files[0])
                else:
                    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                    zip_path = os.path.join(out_dir, 'chatlab_%s.zip' % stamp)
                    with _zipfile.ZipFile(zip_path, 'w', _zipfile.ZIP_DEFLATED) as zf:
                        for p in files:
                            zf.write(p, arcname=os.path.basename(p))
                    fname = os.path.basename(zip_path)
                note = '导入方式：ChatLab → 导入聊天记录 → 选择该文件。'
                if keyword:
                    note = '已按关键词过滤：%s。' % keyword + note
                push.done({'msg_count': total,
                           'download_url': '/api/export/download/' + fname,
                           'filename': fname,
                           'note': note})
                return

            push('export', '正在导出 %d 个会话…' % len(targets), 0.25)

            def _progress(pct, msg):
                push('export', msg, 0.25 + 0.7 * (pct / 100.0))

            if len(targets) == 1 and fmt not in ('xlsx', 'csv', 'json', 'pgsql', 'sql'):
                count, filepath = _export_one(
                    targets[0], out_dir,
                    start_ts=start_ts, end_ts=end_ts, keyword=keyword or None,
                    print_fn=_print, fmt=fmt,
                    display_name=display_name_hint or targets[0].get('display_name'),
                    include_asr=bool(data.get('include_asr')),
                    privacy=privacy)
            else:
                count, filepath, _details = export_chats(
                    targets, out_dir,
                    start_ts=start_ts, end_ts=end_ts, keyword=keyword or None,
                    print_fn=_print, progress_fn=_progress, fmt=fmt,
                    include_asr=bool(data.get('include_asr')),
                    merge=bool(merge),
                    privacy=privacy)

            if count and filepath:
                fname = os.path.basename(filepath)
                note = ''
                if keyword:
                    note = '已按关键词过滤：%s' % keyword
                if len(targets) > 1 and fmt in ('xlsx', 'csv', 'html', 'json', 'pgsql', 'sql', 'md', 'markdown'):
                    note = (note + '；' if note else '') + '多会话已按时间平铺合并'
                if privacy:
                    note = (note + '；' if note else '') + '已启用隐私脱敏（昵称/手机号/身份证）'
                push.done({'msg_count': count,
                           'download_url': '/api/export/download/' + fname,
                           'filename': fname,
                           'note': note or None})
            else:
                push.done({'msg_count': count or 0, 'file': filepath,
                           'note': '没有匹配的消息（请检查日期/关键词）'})
        except Exception as e:
            push.error(str(e))

    threading.Thread(target=_run, daemon=True).start()
    return sse_response(gen)


@export_bp.route('/wordcloud/view/<path:filename>')
def view_wordcloud(filename):
    """GET /api/export/wordcloud/view/<filename> — Serve a generated wordcloud HTML."""
    out_dir = os.path.join(_DATA_ROOT, 'export', 'wordcloud')
    return send_from_directory(os.path.normpath(out_dir), filename)


@export_bp.route('/download/<path:filename>')
def download_file(filename):
    """GET /api/export/download/<filename> — Serve an exported file."""
    out_dir = os.path.join(_DATA_ROOT, 'export')
    return send_from_directory(os.path.normpath(out_dir), filename, as_attachment=True)


@export_bp.route('/wordcloud', methods=['POST'])
def export_wordcloud():
    """POST /api/export/wordcloud — Generate word cloud."""
    data = request.get_json(silent=True) or {}
    chat = data.get('chat', '')
    try:
        year = int(data.get('year', datetime.now().year))
    except (TypeError, ValueError):
        year = datetime.now().year

    decrypted_dir = _decrypted_dir()
    push, gen = create_sse_progress()

    def _run():
        try:
            from wordcloud_gen import generate_wordcloud, extract_text_messages
            from chat_list import scan_chats
            from engine.constants import TZ

            push('analyze', '正在扫描聊天数据...', 0.1)

            chat_info = None
            if chat:
                all_chats, _, _ = scan_chats(decrypted_dir)
                if not all_chats:
                    # Flat backup layout (no contact/session subdirs) —
                    # scan Name2Id from message DBs directly
                    all_chats = _scan_chats_flat(decrypted_dir, chat)
                # Exact match first
                for c in all_chats:
                    if c['username'] == chat or c['display_name'] == chat:
                        chat_info = c
                        break
                # Fuzzy match: collect all matches, let user pick
                if not chat_info:
                    chat_lower = chat.lower()
                    fuzzy_matches = []
                    for c in all_chats:
                        if (chat_lower in c['display_name'].lower()
                                or chat_lower in c['username'].lower()):
                            fuzzy_matches.append({
                                'username': c['username'],
                                'display_name': c['display_name'],
                                'msg_count': c.get('msg_count', 0),
                            })
                    if not fuzzy_matches:
                        push.error(f'未找到匹配的聊天: {chat}')
                        return
                    if len(fuzzy_matches) == 1:
                        # Single match — auto-select
                        m = fuzzy_matches[0]
                        for c in all_chats:
                            if c['username'] == m['username']:
                                chat_info = c
                                break
                    else:
                        push.select(fuzzy_matches)
                        return

            start_ts = int(datetime(year, 1, 1, tzinfo=TZ).timestamp())
            end_ts = int(datetime(year + 1, 1, 1, tzinfo=TZ).timestamp()) - 1

            push('analyze', f'正在提取 {year} 年文本...', 0.3)
            text, msg_count = extract_text_messages(decrypted_dir, chat_info=chat_info,
                                                     start_ts=start_ts, end_ts=end_ts)

            if msg_count == 0:
                # Year filter may exclude all data — retry without time filter
                push('analyze', f'{year} 年无数据，正在全时段提取...', 0.35)
                text, msg_count = extract_text_messages(decrypted_dir, chat_info=chat_info,
                                                         start_ts=None, end_ts=None)
                year_ts = (None, None)
            else:
                year_ts = (start_ts, end_ts)

            push('analyze', f'提取到 {msg_count} 条消息', 0.5)

            out_dir = os.path.join(_DATA_ROOT, 'export', 'wordcloud')
            os.makedirs(out_dir, exist_ok=True)
            ts = datetime.now().strftime('%Y%m%d_%H%M%S')
            out_file = os.path.join(out_dir, f'wordcloud_{ts}.html')
            def _print(msg):
                push('analyze', str(msg), 0.6)

            result = generate_wordcloud(decrypted_dir, chat_info=chat_info,
                                        out_path=out_file, start_ts=year_ts[0], end_ts=year_ts[1],
                                        print_fn=_print)

            if result:
                fname = os.path.basename(result)
                view_url = f'/api/export/wordcloud/view/{fname}'
                push.done({'msg_count': msg_count, 'view_url': view_url, 'filename': fname})
            else:
                push.error('词云生成失败，未找到足够的文本消息')
        except Exception as e:
            push.error(str(e))

    threading.Thread(target=_run, daemon=True).start()
    return sse_response(gen)


@export_bp.route('/report/view/<path:filename>')
def view_report(filename):
    """GET /api/export/report/view/<filename> — Serve a generated report HTML."""
    out_dir = os.path.join(_DATA_ROOT, 'export', 'report')
    return send_from_directory(os.path.normpath(out_dir), filename)


@export_bp.route('/report', methods=['POST'])
def export_report():
    """POST /api/export/report — Generate HTML report."""
    data = request.get_json(silent=True) or {}
    decrypted_dir = _decrypted_dir()
    year = data.get('year')
    start_ts = end_ts = None
    try:
        if year not in (None, '', 0, '0'):
            y = int(year)
            from datetime import datetime
            start_ts = int(datetime(y, 1, 1, 0, 0, 0).timestamp())
            end_ts = int(datetime(y, 12, 31, 23, 59, 59).timestamp())
    except (TypeError, ValueError):
        start_ts = end_ts = None

    push, gen = create_sse_progress()

    def _run():
        try:
            from report_gen import generate_report
            scope = f'{year} 年' if start_ts is not None else '全部时间'
            push('analyze', f'正在收集统计数据（{scope}）... ({decrypted_dir})', 0.2)
            out_file = generate_report(
                decrypted_dir, start_ts=start_ts, end_ts=end_ts)
            if out_file:
                # Copy to web-accessible export dir and return view URL
                fname = os.path.basename(out_file)
                web_dir = os.path.join(_DATA_ROOT, 'export', 'report')
                os.makedirs(web_dir, exist_ok=True)
                import shutil
                shutil.copy2(out_file, os.path.join(web_dir, fname))
                view_url = f'/api/export/report/view/{fname}'
                push.done({'msg_count': 0, 'view_url': view_url, 'filename': fname})
            else:
                push.error('报告生成失败 — 解密数据目录中未找到任何消息。请确认已执行备份与解密。')
        except Exception as e:
            push.error(str(e))

    threading.Thread(target=_run, daemon=True).start()
    return sse_response(gen)


@export_bp.route('/employee', methods=['POST'])
def export_employee():
    """POST /api/export/employee — Export employee chat reports.

    This is a synchronous JSON endpoint (not SSE) because the file upload
    via multipart/form-data makes SSE streaming impractical. All work runs
    in the request thread.
    """
    import tempfile

    excel_path = None
    tmp = None
    if 'file' in request.files:
        f = request.files['file']
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        f.save(tmp.name)
        excel_path = tmp.name
        tmp.close()
    elif request.form.get('excel'):
        excel_path = request.form['excel']

    if not excel_path:
        return jsonify({'success': False, 'error': '请上传员工 Excel 文件'}), 400

    decrypted_dir = _decrypted_dir()
    out_dir = os.path.join(_DATA_ROOT, 'export')
    os.makedirs(out_dir, exist_ok=True)

    try:
        from employee_match import run_employee_export
        import io
        import sys as _sys
        with _stdout_lock:
            old_stdout = _sys.stdout
            _sys.stdout = io.StringIO()
            try:
                run_employee_export(decrypted_dir, excel_path, out_dir)
                output = _sys.stdout.getvalue()
            finally:
                _sys.stdout = old_stdout
        return jsonify({'success': True, 'output': output})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if tmp is not None:
            try:
                os.unlink(tmp.name)
            except OSError:
                pass


_SEARCH_EXPORT_MAX = 10000
_SEARCH_EXPORT_FMTS = ('xlsx', 'html', 'txt', 'csv')


def _write_search_export(hits, fmt, out_path, query=''):
    """Write search hits to xlsx/html/txt/csv. Returns row count."""
    from engine.constants import TZ

    def _ts(v):
        try:
            return datetime.fromtimestamp(int(v), tz=TZ).strftime('%Y-%m-%d %H:%M:%S')
        except (TypeError, ValueError, OSError):
            return ''

    rows = []
    for h in hits:
        rows.append({
            'time': _ts(h.get('create_time')),
            'chat': h.get('chat_display_name') or h.get('chat_id') or '',
            'sender': h.get('sender_display_name') or h.get('sender_username') or '',
            'type': h.get('type_label') or '',
            'text': h.get('text') if h.get('text') is not None else (h.get('snippet') or ''),
            'is_group': '是' if h.get('is_group') else '否',
        })

    if fmt == 'xlsx':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError as e:
            raise RuntimeError('导出 xlsx 需要 openpyxl：pip install openpyxl') from e
        wb = Workbook()
        ws = wb.active
        ws.title = '搜索结果'
        headers = ['时间', '会话', '发送者', '类型', '内容', '群聊']
        ws.append(headers)
        head_fill = PatternFill('solid', fgColor='FFE0EF')
        for col, _ in enumerate(headers, 1):
            cell = ws.cell(1, col)
            cell.font = Font(bold=True)
            cell.fill = head_fill
        for r in rows:
            ws.append([r['time'], r['chat'], r['sender'], r['type'], r['text'], r['is_group']])
            ws.cell(ws.max_row, 5).alignment = Alignment(wrap_text=True)
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 22
        ws.column_dimensions['C'].width = 16
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 60
        ws.column_dimensions['F'].width = 6
        wb.save(out_path)
    elif fmt == 'csv':
        import csv
        with open(out_path, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f)
            w.writerow(['时间', '会话', '发送者', '类型', '内容', '群聊'])
            for r in rows:
                w.writerow([r['time'], r['chat'], r['sender'], r['type'], r['text'], r['is_group']])
    elif fmt == 'txt':
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write('搜索导出 · 查询：%s · 共 %d 条\n' % (query, len(rows)))
            f.write('=' * 60 + '\n\n')
            for r in rows:
                f.write('[%s] %s | %s | %s\n%s\n\n' % (
                    r['time'], r['chat'], r['sender'], r['type'], r['text']))
    else:  # html
        import html as _html
        parts = [
            '<!DOCTYPE html><html><head><meta charset="utf-8">',
            '<title>搜索结果导出</title>',
            '<style>',
            'body{font-family:"Segoe UI",sans-serif;background:#fff5f8;color:#1a1a1a;',
            'margin:0;padding:24px}',
            'h1{font-size:20px;margin:0 0 8px}',
            '.meta{color:#666;font-size:13px;margin-bottom:20px}',
            'table{border-collapse:collapse;width:100%;background:#fff;',
            'border:2px solid #f5a0c0}',
            'th{background:#ffe0ef;text-align:left;padding:8px 10px;',
            'border-bottom:2px solid #f5a0c0;font-size:13px}',
            'td{padding:8px 10px;border-bottom:1px solid #f7d0e0;',
            'vertical-align:top;font-size:13px}',
            'td.content{white-space:pre-wrap;word-break:break-word;max-width:520px}',
            '</style></head><body>',
            '<h1>搜索结果导出</h1>',
            '<div class="meta">查询：%s · 共 %d 条</div>' % (
                _html.escape(query or ''), len(rows)),
            '<table><thead><tr>',
            '<th>时间</th><th>会话</th><th>发送者</th><th>类型</th><th>内容</th>',
            '</tr></thead><tbody>',
        ]
        for r in rows:
            parts.append(
                '<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td>'
                '<td class="content">%s</td></tr>' % (
                    _html.escape(r['time']),
                    _html.escape(r['chat']),
                    _html.escape(r['sender']),
                    _html.escape(r['type']),
                    _html.escape(r['text'])))
        parts.append('</tbody></table></body></html>')
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(''.join(parts))
    return len(rows)


@export_bp.route('/search', methods=['POST'])
def export_search():
    """POST /api/export/search — 将当前搜索命中导出为 Excel/HTML/TXT/CSV。

    Body: {q, format?}  format ∈ xlsx|html|txt|csv，默认 xlsx。
    """
    data = request.get_json(silent=True) or {}
    q = (data.get('q') or '').strip()
    fmt = (data.get('format') or 'xlsx').lower().strip()
    if fmt not in _SEARCH_EXPORT_FMTS:
        return jsonify({'error': '不支持的格式：%s（可选 %s）'
                        % (fmt, '/'.join(_SEARCH_EXPORT_FMTS))}), 400
    if not q:
        return jsonify({'error': '缺少查询参数 q'}), 400

    decrypted_dir = _decrypted_dir()
    own_wxid = current_app.config.get('OWN_WXID', '') or ''
    push, gen = create_sse_progress()

    def _run():
        try:
            from engine.services.search import (
                search_messages, SearchIndexError, MAX_RETAINED_ROWS)

            push('export', '正在检索命中消息…', 0.15)
            cap = min(_SEARCH_EXPORT_MAX, MAX_RETAINED_ROWS)
            try:
                body = search_messages(
                    decrypted_dir, q,
                    page=1, per_page=cap, sort='time_desc',
                    own_wxid=own_wxid or None,
                    max_retained=cap,
                    include_full_text=True,
                    export_all=True)
            except SearchIndexError as e:
                push.error(e.as_dict().get('error') or str(e))
                return
            except ValueError as e:
                push.error(str(e))
                return

            hits = body.get('results') or []
            if not hits:
                push.done({'msg_count': 0, 'note': '没有匹配的消息可导出'})
                return

            push('export', '正在写入 %d 条…' % len(hits), 0.55)
            out_dir = os.path.join(_DATA_ROOT, 'export')
            os.makedirs(out_dir, exist_ok=True)
            stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            fname = 'search_%s.%s' % (stamp, fmt)
            out_path = os.path.join(out_dir, fname)
            count = _write_search_export(hits, fmt, out_path, query=q)

            note = None
            if body.get('truncated'):
                note = ('命中共 %d 条，本次仅导出前 %d 条（内存保护上限）。'
                        '请缩小筛选条件后重试。'
                        % (body.get('total') or count, count))
            push.done({
                'msg_count': count,
                'total_hits': body.get('total'),
                'truncated': bool(body.get('truncated')),
                'download_url': '/api/export/download/' + fname,
                'filename': fname,
                'note': note,
            })
        except Exception as e:
            push.error(str(e))

    threading.Thread(target=_run, daemon=True).start()
    return sse_response(gen)
